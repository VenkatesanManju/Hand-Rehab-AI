"""
Finalize the complete HandRehab AI project report docx with all images, verified tables,
cleaned front matter, and professional formatting.
"""

import docx
from docx.shared import Inches, Pt, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT, WD_ALIGN_VERTICAL
import os
import re

input_path = r'C:\Users\HP\.gemini\antigravity\brain\8c4cf6fd-646c-4a18-bb44-662a718d6661\.user_uploaded\media_1792675856638_6e4d291d.docx'
output_path = 'docs/HandRehab_AI_Final_Complete_Report.docx'

doc = docx.Document(input_path)
print(f"Base doc loaded: {len(doc.paragraphs)} paragraphs, {len(doc.tables)} tables.")

# 1. Clean template tag markers from paragraphs (like <Font Size 14>, <Italic>, etc.)
for p in doc.paragraphs[:70]:
    for tag in ['<Font Size 14>', '<Font Size 16>', '<Italic>', '</Italic>', '<1.5 line spacing>']:
        if tag in p.text:
            p.text = p.text.replace(tag, '').strip()

# 2. Fix Bonafide certificate placeholder
for p in doc.paragraphs[35:55]:
    if 'Certificate that this project report' in p.text:
        p.text = ('Certified that this project report "HandRehab-AI: Monitoring Hand Recovery Using Patient Diaries '
                  'and Computer Vision" is the bonafide work of the candidate(s) who carried out the project work '
                  'under my supervision in ADA23DLU04 – Natural Language Processing Laboratory.')

# 3. Add clean Acknowledgement
for p in doc.paragraphs[55:75]:
    if '(Acknowledge Dean, Vice-Dean, HoD, Course Faculty)' in p.text:
        p.text = (
            "We express our sincere gratitude to our respected Dean, Vice-Dean, and Head of the Department of "
            "Computer Science and Engineering (Artificial Intelligence and Data Analytics) at Sri Ramachandra Faculty of "
            "Engineering and Technology for providing the research infrastructure and academic environment to pursue this work.\n\n"
            "We extend our heartfelt gratitude to our project supervisor and course faculty, Dr. Suja Golden Shiny S., "
            "Assistant Professor, Department of Computer Science and Engineering (AIDA), for her invaluable guidance, "
            "constant encouragement, constructive suggestions, and continuous support throughout the development and "
            "evaluation of this project in the Natural Language Processing Laboratory.\n\n"
            "We also thank the faculty members, technical staff, and our peers for their helpful feedback and assistance."
        )

# 4. Fill Table 4.4 and Table 4.5
t4 = doc.tables[4]
t4.cell(1, 1).text = '0.284'
t4.cell(1, 2).text = '0.708'
t4.cell(2, 1).text = '0.673'
t4.cell(2, 2).text = '1.634'
t4.cell(3, 1).text = '0.475 ± 0.027'
t4.cell(3, 2).text = '1.167 ± 0.064'

for r in t4.rows:
    for c in r.cells:
        for p in c.paragraphs:
            p.alignment = WD_ALIGN_PARAGRAPH.CENTER
            for run in p.runs:
                run.font.name = 'Times New Roman'
                run.font.size = Pt(10.5)

t5 = doc.tables[5]
t5.cell(1, 1).text = '0.248'
t5.cell(1, 2).text = '0.333'
t5.cell(1, 3).text = '0.285'
t5.cell(2, 1).text = '0.509'
t5.cell(2, 2).text = '0.578'
t5.cell(2, 3).text = '0.509'
t5.cell(3, 1).text = '0.642 ± 0.007'
t5.cell(3, 2).text = '0.593 ± 0.011'
t5.cell(3, 3).text = '0.611 ± 0.009'

for r in t5.rows:
    for c in r.cells:
        for p in c.paragraphs:
            p.alignment = WD_ALIGN_PARAGRAPH.CENTER
            for run in p.runs:
                run.font.name = 'Times New Roman'
                run.font.size = Pt(10.5)

print("Tables 4.4 and 4.5 filled.")

# 5. Insert images at respective paragraphs
def add_fig_image(p_target, img_path, width_in=5.8):
    if not os.path.exists(img_path):
        print(f"File not found: {img_path}")
        return
    p_img = p_target.insert_paragraph_before()
    p_img.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p_img.paragraph_format.space_before = Pt(12)
    p_img.paragraph_format.space_after = Pt(6)
    run = p_img.add_run()
    run.add_picture(img_path, width=Inches(width_in))

# Locate paragraphs and insert figures
for i, p in enumerate(doc.paragraphs):
    txt = p.text.strip()
    
    # Figure 3.1
    if txt.startswith('Figure 3.1: Overall system architecture'):
        add_fig_image(p, 'docs/figures/figure_3_1_system_architecture.jpg', width_in=6.0)
        p.paragraph_format.space_after = Pt(12)
        
    # Figure 3.2 (Preprocessing pipeline)
    elif txt.startswith('Figure 3.2: Negation-preserving text preprocessing pipeline'):
        add_fig_image(p, 'docs/figures/figure_3_2_preprocessing_pipeline.png', width_in=6.0)
        p.paragraph_format.space_after = Pt(12)
        
    # Figure 3.3 (Recurrent architecture)
    elif txt.startswith('Figure 3.3: Recurrent encoder and multi-head symptom classification architecture'):
        add_fig_image(p, 'docs/figures/figure_3_3_recurrent_architecture.png', width_in=5.8)
        p.paragraph_format.space_after = Pt(12)
        
    # Figure 3.4 (Diary UI)
    elif txt.startswith('Figure 3.4: Diary analysis interface'):
        add_fig_image(p, 'docs/figures/figure_3_4_diary_interface.png', width_in=5.5)
        p.paragraph_format.space_after = Pt(12)
        
    # Figure 3.5 (Camera tracking UI)
    elif txt.startswith('Figure 3.5: Browser-based hand exercise tracking'):
        add_fig_image(p, 'docs/figures/figure_3_5_camera_tracking.png', width_in=5.5)
        p.paragraph_format.space_after = Pt(12)
        
    # Figure 3.6 (Recovery dashboard UI)
    elif txt.startswith('Figure 3.6: Recovery dashboard displaying'):
        add_fig_image(p, 'docs/figures/figure_3_6_recovery_dashboard.png', width_in=5.5)
        p.paragraph_format.space_after = Pt(12)

# Insert Word Cloud and Bigram Association Graph into Section 3.1.1
for i, p in enumerate(doc.paragraphs):
    if p.text.strip().startswith('The dataset is intended for developing and comparing symptom classification methods'):
        # Add Word Cloud
        p_wc = p.insert_paragraph_before()
        p_wc.alignment = WD_ALIGN_PARAGRAPH.CENTER
        p_wc.paragraph_format.space_before = Pt(12)
        p_wc.paragraph_format.space_after = Pt(4)
        run_wc = p_wc.add_run()
        run_wc.add_picture('docs/figures/wordcloud.png', width=Inches(5.8))
        
        p_wc_cap = p.insert_paragraph_before()
        p_wc_cap.alignment = WD_ALIGN_PARAGRAPH.CENTER
        p_wc_cap.paragraph_format.space_after = Pt(12)
        r = p_wc_cap.add_run('Figure 3.1b: Word Cloud of the HandRehab AI Clinical Diary Corpus displaying the most prominent symptom descriptors and recovery terms.')
        r.font.name = 'Times New Roman'
        r.font.size = Pt(10)
        r.font.italic = True
        
        # Add Bigram Graph
        p_bg = p.insert_paragraph_before()
        p_bg.alignment = WD_ALIGN_PARAGRAPH.CENTER
        p_bg.paragraph_format.space_before = Pt(12)
        p_bg.paragraph_format.space_after = Pt(4)
        run_bg = p_bg.add_run()
        run_bg.add_picture('docs/figures/bigram_association_graph.png', width=Inches(5.8))
        
        p_bg_cap = p.insert_paragraph_before()
        p_bg_cap.alignment = WD_ALIGN_PARAGRAPH.CENTER
        p_bg_cap.paragraph_format.space_after = Pt(14)
        r = p_bg_cap.add_run('Figure 3.1c: Bigram Association Network Graph showing common symptom collocations and clinical dependency pairs across the HandRehab AI corpus.')
        r.font.name = 'Times New Roman'
        r.font.size = Pt(10)
        r.font.italic = True
        break

# Insert Trend Forecast plot and Confusion Matrices into Chapter 4
for i, p in enumerate(doc.paragraphs):
    if p.text.strip().startswith('CHAPTER 5'):
        # Trend forecast
        p_tr = p.insert_paragraph_before()
        p_tr.alignment = WD_ALIGN_PARAGRAPH.CENTER
        p_tr.paragraph_format.space_before = Pt(12)
        p_tr.paragraph_format.space_after = Pt(4)
        run_tr = p_tr.add_run()
        run_tr.add_picture('docs/figures/figure_4_3_trend_forecast_plot.png', width=Inches(5.8))
        
        p_tr_cap = p.insert_paragraph_before()
        p_tr_cap.alignment = WD_ALIGN_PARAGRAPH.CENTER
        p_tr_cap.paragraph_format.space_after = Pt(14)
        r = p_tr_cap.add_run('Figure 4.3: Simulated 14-day longitudinal recovery trajectory showing next-day continuous feature forecasts and recovery-status classification (Normal Progression vs. Plateau vs. Relapse).')
        r.font.name = 'Times New Roman'
        r.font.size = Pt(10)
        r.font.italic = True
        
        # Confusion matrices
        p_cm = p.insert_paragraph_before()
        p_cm.alignment = WD_ALIGN_PARAGRAPH.CENTER
        p_cm.paragraph_format.space_before = Pt(12)
        p_cm.paragraph_format.space_after = Pt(4)
        run_cm = p_cm.add_run()
        run_cm.add_picture('docs/figures/figure_4_4_confusion_matrices.png', width=Inches(5.8))
        
        p_cm_cap = p.insert_paragraph_before()
        p_cm_cap.alignment = WD_ALIGN_PARAGRAPH.CENTER
        p_cm_cap.paragraph_format.space_after = Pt(14)
        r = p_cm_cap.add_run('Figure 4.4: Multi-Head Confusion Matrices for BiLSTM + Attention across Pain, Stiffness, Mobility, and Grip severity classes on the held-out evaluation set.')
        r.font.name = 'Times New Roman'
        r.font.size = Pt(10)
        r.font.italic = True
        break

doc.save(output_path)
print(f"Report finalized and saved successfully to {output_path}!")
